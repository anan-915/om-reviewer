#!/usr/bin/env python3
"""Build an English author-facing DOCX from om-reviewer JSON."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

from validate_review_output import validate_payload


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = SCRIPT_DIR.parent / "assets" / "review-report-template.docx"


def _set_run_font(run: Any, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")


def _configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    for style_name in ("Normal", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")
        rfonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    list_style = doc.styles["List Number"]
    list_style.paragraph_format.left_indent = Inches(0.25)
    list_style.paragraph_format.first_line_indent = Inches(-0.25)

    props = doc.core_properties
    props.title = "Comments to the Authors"
    props.subject = "Peer review report"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = ""
    props.comments = ""


def _clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def create_template(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _configure_document(doc)
    _clear_body(doc)
    doc.add_paragraph("")
    doc.save(path)


def _default_assessment(recommendation: str) -> str:
    return {
        "Reject": "I recommend rejection.",
        "Major Revision": "I recommend major revision.",
        "Minor Revision": "I recommend minor revision.",
        "Accept": "I recommend acceptance.",
    }[recommendation]


def _append_paragraph(doc: Document, text: str, first_line: bool = False) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Inches(0.33)
    run = paragraph.add_run(text.strip())
    _set_run_font(run)


def _concern_body(concern: dict[str, Any]) -> str:
    return str(concern.get("comment", "")).strip()


def _append_word_equation(paragraph: Any, linear_formula: str) -> None:
    separator = paragraph.add_run()
    _set_run_font(separator)
    separator.add_break(WD_BREAK.LINE)

    equation = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_properties = OxmlElement("m:rPr")
    math_style = OxmlElement("m:sty")
    math_style.set(qn("m:val"), "p")
    math_properties.append(math_style)
    math_text = OxmlElement("m:t")
    math_text.set(qn("xml:space"), "preserve")
    math_text.text = linear_formula.strip()
    math_run.append(math_properties)
    math_run.append(math_text)
    equation.append(math_run)
    paragraph._p.append(equation)


def build_docx(payload: dict[str, Any], output: Path, template: Path) -> None:
    doc = Document(template) if template.exists() else Document()
    _configure_document(doc)
    _clear_body(doc)

    _append_paragraph(doc, str(payload["manuscript_summary"]), first_line=True)
    assessment = str(payload.get("author_overall_assessment", "")).strip()
    if not assessment:
        assessment = _default_assessment(str(payload["recommendation"]))
    _append_paragraph(doc, assessment, first_line=True)

    for concern in payload.get("concerns", []):
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True

        body = paragraph.add_run(_concern_body(concern))
        _set_run_font(body, bold=False)
        for equation in concern.get("equations", []):
            _append_word_equation(paragraph, str(equation))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    expected_equations = sum(len(concern.get("equations", [])) for concern in payload.get("concerns", []))
    audit_docx(output, expected_concerns=len(payload.get("concerns", [])), expected_equations=expected_equations)


def audit_docx(path: Path, expected_concerns: int | None, expected_equations: int | None = None) -> None:
    doc = Document(path)
    section = doc.sections[0]
    checks = {
        "A4 width": abs(section.page_width.mm - 210) < 0.2,
        "A4 height": abs(section.page_height.mm - 297) < 0.2,
        "top margin": abs(section.top_margin.inches - 1) < 0.02,
        "bottom margin": abs(section.bottom_margin.inches - 1) < 0.02,
        "left margin": abs(section.left_margin.inches - 1.25) < 0.02,
        "right margin": abs(section.right_margin.inches - 1.25) < 0.02,
        "no tables": len(doc.tables) == 0,
        "paragraph count": expected_concerns is None or len(doc.paragraphs) == 2 + expected_concerns,
        "numbered concerns": expected_concerns is None
        or sum(p.style.name == "List Number" for p in doc.paragraphs) == expected_concerns,
        "native equations": expected_equations is None
        or len(doc._element.xpath(".//m:oMath")) == expected_equations,
    }
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise ValueError(f"DOCX structural audit failed: {', '.join(failed)}")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text and run.font.name not in (None, "Times New Roman"):
                raise ValueError(f"Unexpected font in output: {run.font.name}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", nargs="?", type=Path, help="UTF-8 review JSON")
    parser.add_argument("output", nargs="?", type=Path, help="Output DOCX")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--make-template", type=Path, help="Create a blank review template and exit")
    args = parser.parse_args()

    if args.make_template:
        create_template(args.make_template)
        audit_docx(args.make_template, expected_concerns=None, expected_equations=0)
        print(f"Created template: {args.make_template}")
        return 0

    if args.input is None or args.output is None:
        parser.error("input and output are required unless --make-template is used")

    try:
        payload = json.loads(args.input.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        print(f"ERROR: Could not read review JSON: {exc}", file=sys.stderr)
        return 1
    if not isinstance(payload, dict):
        print("ERROR: Top-level JSON value must be an object", file=sys.stderr)
        return 1

    errors, warnings = validate_payload(payload)
    for warning in warnings:
        print(f"WARNING: {warning}")
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        return 1

    if not args.template.exists():
        create_template(args.template)
    build_docx(payload, args.output, args.template)
    print(f"Created review DOCX: {args.output}")
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    raise SystemExit(main())
